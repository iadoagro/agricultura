# -*- coding: utf-8 -*-
"""
Gera js/dados-mecanizacao-historico.js a partir da aba "geral" de
mecanizacao26.xlsx — os exercicios ENCERRADOS (2023, 2024 e 2025).

Uso (na raiz do projeto):
    python tools/gerar_historico_mecanizacao.py

Por que um arquivo separado do de 2026:
  - so 2026 continua recebendo lancamentos. O painel le o historico de um
    arquivo fixo e junta com o pacote do ano corrente, entao publicar uma
    planilha nova de 2026 (aba "Atualizar dados") nao apaga os anos anteriores;
  - a aba "geral" nao tem a mesma estrutura de colunas da aba "dados", e o
    layout das colunas de cultura MUDA de um ano para o outro (ver CULTURAS).

As normalizacoes (escritorio, tecnico, maquinas, implementos)
sao reaproveitadas de gerar_dados_mecanizacao.py para que os dois arquivos
falem a mesma lingua e o consolidado nao conte a mesma coisa duas vezes.
"""
import io
import json
import os
import sys
from datetime import datetime

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import gerar_dados_mecanizacao as base

RAIZ = base.RAIZ
XLSX = base.XLSX
SAIDA = os.path.join(RAIZ, "js", "dados-mecanizacao-historico.js")
ABA = "geral"
ANOS = ("2023", "2024", "2025")

NAO_INF = base.NAO_INF

# Colunas fixas da aba "geral" (indice 0-based; o cabecalho tem espacos e
# duplicatas que impedem a busca por nome usada na aba "dados").
C = {
    "ano": 0, "carimbo": 1, "email": 2, "email2": 3,
    "escritorio": 4, "vistoria": 5, "tecnico": 6,
    # a coluna 9 e o CPF: nao e lida, para que nenhum dado pessoal saia daqui
    "produtor": 7, "sexo": 8, "civil": 10, "assoc": 13, "dap": 14,
    "municipio": 17, "endereco": 18, "propriedade": 19, "ponto": 20,
    "horas": 21, "maquina": 22, "total_mec": 38, "geox": 39, "geoy": 40,
    "trator": 43, "tipo_trator": 44, "tipo_impl": 45, "nome_impl": 46,
    "obs": 51, "formulario": 52, "acudes": 71,
}
DAE_VALORES = [50] + list(range(62, 71))

# Culturas: (coluna do nome, coluna da area, coluna do sistema de cultivo).
# area = None significa "a planilha nao registrou a area dessa cultura".
#
#  - 2023: o formulario tinha uma cultura so, com area propria (col. 24, que
#          bate exatamente com "Total mecanizado" em todas as linhas do ano);
#  - 2024: mesma cultura unica, mas SEM coluna de area preenchida;
#  - 2025: ate tres culturas. A primeira nao tem coluna de area (a col. 24
#          repete o nome da cultura); as adicionais usam as col. 32 e 36.
#
# Para 2024 e para a primeira cultura de 2025 a area vem de "Total mecanizado",
# que e como 2023 e 2026 se comportam quando ha uma unica cultura declarada.
CULTURAS = {
    "2023": [(23, 24, 25)],
    "2024": [(23, None, 25)],
    "2025": [(23, None, 25), (27, 32, 33), (35, 36, 37)],
}


def main():
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    if ABA not in wb.sheetnames:
        raise SystemExit('Aba "%s" nao encontrada em %s' % (ABA, XLSX))
    ws = wb[ABA]

    linhas = ws.iter_rows(values_only=True)
    next(linhas)  # cabecalho

    produtores = set()
    registros = []
    por_ano = {}
    qualidade = {"sem_data_valida": 0, "vistoria_outro_ano": 0,
                 "sem_geo": 0, "sem_formulario": 0, "fora_do_periodo": 0}

    for r in linhas:
        if not any(c is not None and base.texto(c) != "" for c in r):
            continue
        ano = base.texto(r[C["ano"]])[:4]
        if ano not in ANOS:
            # linhas de teste e lancamentos de 2026, que vem da aba "dados"
            qualidade["fora_do_periodo"] += 1
            continue

        def v(k):
            j = C[k]
            return r[j] if j < len(r) else None

        data = base.data_iso(v("carimbo"))
        vistoria = base.data_iso(v("vistoria"))
        if not data:
            data = vistoria
            qualidade["sem_data_valida"] += 1
        if vistoria and vistoria[:4] != data[:4]:
            qualidade["vistoria_outro_ano"] += 1

        total_mec = round(base.numero(v("total_mec")), 2)

        culturas = []
        for i, (ic, ia, isis) in enumerate(CULTURAS[ano]):
            nome = base.texto(r[ic]) if ic < len(r) else ""
            if not nome or base.eh_nao_informado(nome):
                continue
            if ia is None:
                # sem coluna de area: a primeira cultura herda o total mecanizado
                area = total_mec if i == 0 else 0.0
            else:
                area = round(base.numero(r[ia]), 2)
            culturas.append([base.rotulo(nome), area, base.rotulo(r[isis])])

        dae = round(sum(base.numero(r[j]) for j in DAE_VALORES if j < len(r)), 2)

        gx = base.numero(r[C["geox"]] if C["geox"] < len(r) else None)
        gy = base.numero(r[C["geoy"]] if C["geoy"] < len(r) else None)
        if not gx or not gy:
            qualidade["sem_geo"] += 1

        form = base.texto(v("formulario"))
        if not form.startswith("http"):
            form = ""
            qualidade["sem_formulario"] += 1

        produtor = base.texto(v("produtor"))
        # produtores distintos pelo NOME normalizado (antes vinha do CPF)
        if produtor:
            produtores.add(" ".join(base.sem_acento(produtor).lower().split()))
        dap = base.texto(v("dap"))

        registros.append({
            "d": data,        # data de insercao (Carimbo de data/hora)
            "ex": ano,        # exercicio: a coluna "Ano" manda, nao o carimbo
            "dv": vistoria,
            "pc": base.rotulo(v("ponto")),
            "mun": base.rotulo(v("municipio")),
            "esc": base.norm_escritorio(v("escritorio")),
            "rt": base.norm_tecnico(v("tecnico")),
            "alim": base.norm_alimentador(v("email"), v("email2")),
            "prod": base.titulo(produtor) if produtor.isupper() else produtor,
            "sexo": base.rotulo(v("sexo")),
            "ec": base.norm_estado_civil(v("civil")),
            "dap": dap if dap in ("Sim", "Não", "Vencida") else NAO_INF,
            "assoc": base.rotulo(v("assoc")),
            "loc": base.texto(v("endereco")),
            "propr": base.rotulo(v("propriedade")),
            "cult": culturas,
            "ha": total_mec,
            "hrs": round(base.numero(v("horas")), 2),
            "ac": int(base.numero(v("acudes"))),
            "tt": base.rotulo(v("tipo_trator")),
            "maq": base.extrair(v("trator"), base.MAQUINAS) or base.extrair(v("maquina"), base.MAQUINAS),
            "impl": base.extrair(v("tipo_impl"), base.IMPLEMENTOS) or base.extrair(v("nome_impl"), base.IMPLEMENTOS),
            "dae": dae,
            "form": form,
            "obs": base.texto(v("obs")),
        })
        por_ano[ano] = por_ano.get(ano, 0) + 1

    registros.sort(key=lambda x: (x["ex"], x["d"]))

    meta = {
        "arquivo": os.path.basename(XLSX),
        "aba": ABA,
        "gerado_em": datetime.now().strftime("%d/%m/%Y %H:%M"),
        "registros": len(registros),
        "produtores": len(produtores),
        "anos": sorted(por_ano),
        "por_ano": por_ano,
        "qualidade": qualidade,
    }

    corpo = json.dumps({"meta": meta, "registros": registros},
                       ensure_ascii=False, separators=(",", ":"))
    with io.open(SAIDA, "w", encoding="utf-8", newline="\n") as f:
        f.write("/* Historico da mecanizacao e acudagem (2023-2025) — SEAGRI/AC\n")
        f.write("   GERADO AUTOMATICAMENTE por tools/gerar_historico_mecanizacao.py\n")
        f.write("   Fonte: %s (aba \"%s\") — %s\n" % (meta["arquivo"], ABA, meta["gerado_em"]))
        f.write("   Exercicios encerrados: o painel junta este arquivo com os dados do\n")
        f.write("   ano corrente, que continuam vindo da aba \"dados\". Nao editar a mao. */\n")
        f.write("window.DADOS_MECANIZACAO_HISTORICO = ")
        f.write(corpo)
        f.write(";\n")

    print("OK -> %s" % SAIDA)
    print("  registros: %d | produtores distintos: %d" % (meta["registros"], meta["produtores"]))
    print("  por ano: %s" % por_ano)
    print("  qualidade: %s" % qualidade)
    print("  tamanho: %.1f KB" % (os.path.getsize(SAIDA) / 1024.0))


if __name__ == "__main__":
    main()
