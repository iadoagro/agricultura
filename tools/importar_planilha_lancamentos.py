# -*- coding: utf-8 -*-
"""
Importa as fichas da planilha de producao do Google Forms ("Mecanizacao 2026 -
Producao.xlsx": abas "anteriores" = 2023-2025 e "2026") para a tabela
mecanizacao_lancamentos do Supabase, a mesma da aba "Lancamento" do painel.

Uso (na raiz do projeto):
    python tools/importar_planilha_lancamentos.py "C:/caminho/planilha.xlsx" SAIDA_DIR

Gera arquivos SQL em SAIDA_DIR (lotes de 300 linhas) para rodar com
    npx supabase db query --linked --project-ref <ref> -f <arquivo>
Os arquivos tem CPF e data de nascimento: gere fora do repositorio e apague
depois de rodar.

- "Lancado por" (criado_por_email) = coluna "Endereco de e-mail";
  "criado_em" = "Carimbo de data/hora".
- Pode rodar de novo: cada linha tem uma chave_importacao (hash dos campos
  que identificam a ficha) e o insert ignora as que ja existem.
- Normalizacoes (escritorio, estado civil, maquinas, implementos, culturas
  por ano) vem de gerar_dados_mecanizacao.py / gerar_historico_mecanizacao.py.
"""
import hashlib
import json
import os
import re
import sys
from datetime import datetime

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import gerar_dados_mecanizacao as base
from gerar_historico_mecanizacao import CULTURAS as CULTURAS_ANTERIORES

LOTE = 300

MUNICIPIOS = [
    'Acrelândia', 'Assis Brasil', 'Brasiléia', 'Bujari', 'Capixaba', 'Cruzeiro do Sul',
    'Epitaciolândia', 'Feijó', 'Jordão', 'Mâncio Lima', 'Manoel Urbano', 'Marechal Thaumaturgo',
    'Plácido de Castro', 'Porto Acre', 'Porto Walter', 'Rio Branco', 'Rodrigues Alves',
    'Santa Rosa do Purus', 'Sena Madureira', 'Senador Guiomard', 'Tarauacá', 'Xapuri',
]
_MUN = {base.sem_acento(m).lower(): m for m in MUNICIPIOS}

# Nomes das culturas como no formulario do painel (js/lancamento-campos.js).
CULTURAS_FICHA = [
    'Pastagem - Corte', 'Pastagem - Leite', 'Milho - safra', 'Milho - safrinha',
    'Soja - safra', 'Soja - safrinha', 'Feijão', 'Mandioca', 'Açaí', 'Pupunha',
    'Seringueira', 'Banana', 'Café', 'Laranja', 'Limão', 'Maracujá', 'Graviola',
    'Abacaxi', 'Outras',
]
def _chave_cultura(s):
    return re.sub(r'[^a-z0-9]', '', base.sem_acento(s).lower())


_CULT = {_chave_cultura(c): c for c in CULTURAS_FICHA}
_CULT.update({'outros': 'Outras', 'outro': 'Outras', 'outra': 'Outras'})

# Colunas da aba "2026" (0-based). Na aba "anteriores" tudo anda +1 (coluna "Ano").
C = {
    'carimbo': 0, 'email': 1, 'email2': 2, 'escritorio': 3, 'vistoria': 4, 'tecnico': 5,
    'produtor': 6, 'sexo': 7, 'cpf': 8, 'civil': 9, 'nascimento': 10, 'telefone': 11,
    'assoc': 12, 'dap': 13, 'indigena': 14, 'etnia': 15, 'municipio': 16, 'endereco': 17,
    'propriedade': 18, 'ponto': 19, 'horas': 20, 'maquina': 21, 'total_mec': 37,
    'geox': 38, 'geoy': 39, 'zona': 40, 'patrimonio': 41, 'trator': 42, 'tipo_uso': 43,
    'tipo_impl': 44, 'nome_impl': 45, 'obs': 50, 'formulario': 51, 'acudes': 70,
}
DAE_NUMEROS = [47] + list(range(52, 61))
DAE_VALORES = [49] + list(range(61, 70))
CULTURAS_2026 = [(22, 23, 24), (26, 27, 28), (30, 31, 32), (34, 35, 36)]


def nulo(s):
    s = base.texto(s)
    return None if base.eh_nao_informado(s) else s


def data(v):
    """Data ISO valida (1900-2099) ou None — a planilha tem datas digitadas
    impossiveis ("11/28/0023") que o Postgres recusa."""
    iso = base.data_iso(v)
    try:
        d = datetime.strptime(iso, '%Y-%m-%d')
    except ValueError:
        return None
    return iso if 1900 <= d.year <= 2099 else None


def municipio(v):
    return _MUN.get(base.sem_acento(base.texto(v)).lower(), nulo(v))


def escritorio(v):
    e = base.norm_escritorio(v)
    if e == base.NAO_INF:
        return None
    if e == 'Transacreana':
        return 'Escritório Local da Transacreana'
    return 'Escritório Local de ' + _MUN.get(base.sem_acento(e).lower(), e)


def cultura(v):
    t = base.texto(v)
    return _CULT.get(_chave_cultura(t), t)


def sistema(v):
    c = base.sem_acento(base.texto(v)).lower()
    if not c or base.eh_nao_informado(c):
        return None
    if 'ilpf' in c or 'floresta' in c:
        return 'ILPF'
    if 'ilp' in c or 'pecuaria' in c:
        return 'ILP'
    if 'saf' in c or 'agroflorestal' in c:
        return 'SAF'
    if 'rocado' in c or c.endswith(' ro') or c == 'ro':
        return 'RO'
    if 'mono' in c or c.startswith('m'):
        return 'M'
    return None


def numero_ou_nulo(v):
    n = base.numero(v)
    return round(n, 2) if n else None


def ler_aba(ws, off):
    linhas = ws.iter_rows(min_row=2, values_only=True)
    for r in linhas:
        r = list(r)
        if not base.texto(r[off + C['produtor']]) and not base.texto(r[off + C['email']]):
            continue
        g = lambda k: r[off + C[k]] if off + C[k] < len(r) else None
        ano = base.texto(r[0])[:4] if off else None
        if off and ano not in CULTURAS_ANTERIORES:
            continue  # linhas de teste/sem ano na aba "anteriores"
        yield r, g, ano


def montar(r, g, ano, off):
    total_mec = numero_ou_nulo(g('total_mec'))
    culturas = []
    layout = [(a + 0, b, c) for a, b, c in CULTURAS_ANTERIORES[ano]] if ano else \
        [(a + off, b + off, c + off) for a, b, c in CULTURAS_2026]
    for i, (ic, ia, isis) in enumerate(layout):
        nome = base.texto(r[ic]) if ic < len(r) else ''
        if not nome or base.eh_nao_informado(nome) or _chave_cultura(nome) == 'naoseaplica':
            continue
        area =(total_mec if i == 0 else None) if ia is None else numero_ou_nulo(r[ia] if ia < len(r) else None)
        culturas.append({'cultura': cultura(nome), 'area_ha': area, 'sistema_cultivo': sistema(r[isis] if isis < len(r) else None)})

    daes = []
    for jn, jv in zip(DAE_NUMEROS, DAE_VALORES):
        num = base.texto(r[jn + off]) if jn + off < len(r) else ''
        val = base.numero(r[jv + off]) if jv + off < len(r) else 0
        if (num and not re.fullmatch(r'0+', num)) or val:
            daes.append({'numero': num if num and not re.fullmatch(r'0+', num) else None, 'valor': round(val, 2) if val else None})

    gx, gy = base.texto(g('geox')), base.texto(g('geoy'))
    zona = base.texto(g('zona'))
    pontos = [{'ponto': 1, 'x': gx, 'y': gy, 'zona': zona if zona in ('18', '19') else ''}] if (base.numero(gx) and base.numero(gy)) else []

    email = ''
    for v in (g('email'), g('email2')):
        t = base.texto(v).lower()
        if '@' in t:
            email = t
            break

    sexo = base.texto(g('sexo'))[:1].upper()
    dap = base.texto(g('dap'))
    indig = base.texto(g('indigena'))
    uso = base.sem_acento(base.texto(g('tipo_uso'))).lower()
    ponto = base.texto(g('ponto'))
    produtor = base.texto(g('produtor'))
    form = base.texto(g('formulario'))
    carimbo = g('carimbo')
    acudes = int(base.numero(g('acudes')))
    tipo_trator = nulo(g('trator'))

    reg = {
        'criado_em': carimbo.isoformat() + '-05:00' if hasattr(carimbo, 'isoformat') else None,
        'criado_por_email': email or None,
        'tipo_servico': ponto if ponto in ('Mecanização', 'Açudagem') else None,
        'data_vistoria': data(g('vistoria')),
        'escritorio_local': escritorio(g('escritorio')),
        'responsavel_tecnico': None if base.norm_tecnico(g('tecnico')) == base.NAO_INF else base.norm_tecnico(g('tecnico')),
        'nome_beneficiario': (base.titulo(produtor) if produtor.isupper() else produtor) or '(sem nome)',
        'cpf': nulo(g('cpf')),
        'data_nascimento': data(g('nascimento')),
        'estado_civil': None if base.norm_estado_civil(g('civil')) == base.NAO_INF else base.norm_estado_civil(g('civil')),
        'sexo': sexo if sexo in ('M', 'F') else None,
        'indigena': True if indig == 'Sim' else (False if indig == 'Não' else None),
        'etnia': nulo(g('etnia')),
        'possui_dap': dap if dap in ('Sim', 'Não', 'Vencida') else None,
        'associacao_cooperativa': nulo(g('assoc')),
        'telefone': nulo(g('telefone')),
        'endereco': nulo(g('endereco')),
        'nome_propriedade': nulo(g('propriedade')),
        'municipio': municipio(g('municipio')),
        'culturas': culturas[:4],
        'area_total_ha': total_mec,
        'horas_maquina': numero_ou_nulo(g('horas')),
        'quantidade_acudes': acudes or None,
        'pontos_geo': pontos,
        'tipo_trator': tipo_trator,
        'maquinas': base.extrair(g('trator'), base.MAQUINAS) or base.extrair(g('maquina'), base.MAQUINAS),
        'implementos': base.extrair(g('tipo_impl'), base.IMPLEMENTOS) or base.extrair(g('nome_impl'), base.IMPLEMENTOS),
        'tipo_uso': 'Não agrícola' if uso.startswith('nao agr') else ('Agrícola' if uso.startswith('agr') else None),
        'num_identificacao_patrimonio': nulo(g('patrimonio')),
        'daes': daes[:10],
        'observacao': nulo(g('obs')),
        'formulario_url': form if form.startswith('http') else None,
        'origem': 'planilha',
    }
    if not reg['criado_em'] and reg['data_vistoria']:
        reg['criado_em'] = reg['data_vistoria'] + 'T12:00:00-05:00'
    return reg


def main():
    if len(sys.argv) < 3:
        raise SystemExit(__doc__)
    xlsx, saida = sys.argv[1], sys.argv[2]
    os.makedirs(saida, exist_ok=True)
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)

    registros, vistos = [], {}
    for nome_aba, off in (('anteriores', 1), ('2026', 0)):
        if nome_aba not in wb.sheetnames:
            print('aba "%s" nao encontrada, pulando' % nome_aba)
            continue
        n = 0
        for r, g, ano in ler_aba(wb[nome_aba], off):
            reg = montar(r, g, ano, off)
            base_chave = '|'.join(str(reg[k] or '') for k in ('criado_em', 'criado_por_email', 'cpf', 'nome_beneficiario', 'data_vistoria'))
            # fichas identicas repetidas na planilha continuam sendo linhas distintas
            vistos[base_chave] = vistos.get(base_chave, 0) + 1
            reg['chave_importacao'] = hashlib.sha1(('%s#%d' % (base_chave, vistos[base_chave])).encode('utf-8')).hexdigest()
            registros.append(reg)
            n += 1
        print('%s: %d linhas' % (nome_aba, n))

    for i in range(0, len(registros), LOTE):
        lote = json.dumps(registros[i:i + LOTE], ensure_ascii=False)
        cols = ', '.join(registros[0].keys())
        sql = (
            'insert into public.mecanizacao_lancamentos (%s)\n'
            'select %s from jsonb_populate_recordset(null::public.mecanizacao_lancamentos, $imp$%s$imp$::jsonb)\n'
            'on conflict (chave_importacao) do nothing;\n'
        ) % (cols, cols, lote)
        caminho = os.path.join(saida, 'lote_%03d.sql' % (i // LOTE + 1))
        with open(caminho, 'w', encoding='utf-8') as f:
            f.write(sql)
    print('total: %d registros em %d lote(s) -> %s' % (len(registros), (len(registros) + LOTE - 1) // LOTE, saida))


if __name__ == '__main__':
    main()
